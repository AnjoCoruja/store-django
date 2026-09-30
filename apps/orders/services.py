"""Criação de pedidos e planilha (XLSX). Os preços são recalculados no servidor
a partir do banco — o valor enviado pelo navegador nunca é usado."""
from decimal import Decimal
from io import BytesIO

from django.db import transaction
from django.utils import timezone

from apps.products.models import Product

from .models import Order, OrderItem

MAX_ITEMS = 100
MAX_QTY = 10000


class OrderError(ValueError):
    pass


def tier_for(total_qty: int) -> str:
    if total_qty >= 24:
        return "w24"
    if total_qty >= 6:
        return "w6"
    return ""


def unit_price(product: Product, tier: str) -> Decimal:
    base = product.price
    if tier == "w24":
        return product.wholesale_price_24 if product.wholesale_price_24 is not None else max(Decimal("0"), base - 7)
    if tier == "w6":
        return product.wholesale_price_6 if product.wholesale_price_6 is not None else max(Decimal("0"), base - 5)
    return base


def _clean_variant(value, allowed):
    value = (str(value).strip() if value else "")[:60]
    if allowed and value and value not in allowed:
        raise OrderError(f"Opção inválida: {value}")
    return value


@transaction.atomic
def create_order(items: list[dict], customer_name: str = "") -> Order:
    if not isinstance(items, list) or not items:
        raise OrderError("Carrinho vazio.")
    if len(items) > MAX_ITEMS:
        raise OrderError("Itens demais no carrinho.")
    ids = set()
    for it in items:
        try:
            ids.add(int(it.get("id")))
        except (TypeError, ValueError, AttributeError):
            raise OrderError("Item inválido.")
    products = {
        p.pk: p for p in Product.objects.filter(pk__in=ids, is_active=True, is_published=True)
    }
    parsed = []
    for it in items:
        product = products.get(int(it["id"]))
        if product is None:
            raise OrderError("Produto indisponível.")
        try:
            qty = int(it.get("qty", 1))
        except (TypeError, ValueError):
            raise OrderError("Quantidade inválida.")
        if qty < 1 or qty > MAX_QTY:
            raise OrderError("Quantidade inválida.")
        size = _clean_variant(it.get("size"), product.available_sizes())
        color = _clean_variant(it.get("color"), product.available_colors())
        parsed.append((product, qty, size, color))

    total_qty = sum(q for _, q, _, _ in parsed)
    tier = tier_for(total_qty)
    order = Order.objects.create(
        customer_name=(customer_name or "").strip()[:120], total_qty=total_qty, tier=tier
    )
    total = Decimal("0.00")
    for product, qty, size, color in parsed:
        up = unit_price(product, tier)
        sub = up * qty
        total += sub
        OrderItem.objects.create(
            order=order, product=product, product_name=product.name, line=product.line,
            size=size, color=color, qty=qty, unit_price=up, subtotal=sub,
        )
    order.total = total
    order.save(update_fields=["total"])
    return order


TIER_LABELS = {"": "Varejo", "w6": "Atacado 6+ peças", "w24": "Caixa 24+ peças"}


def order_xlsx(order: Order) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Pedido"
    ws.append([f"Red Blue Line — Pedido {order.code}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Data: {timezone.localtime(order.created_at):%d/%m/%Y %H:%M}", "", f"Faixa: {TIER_LABELS.get(order.tier, '')}"])
    if order.customer_name:
        ws.append([f"Cliente: {order.customer_name}"])
    ws.append([])
    header = ["Produto", "Linha", "Cor", "Tamanho", "Qtd", "Preço unit. (R$)", "Subtotal (R$)"]
    ws.append(header)
    hdr_row = ws.max_row
    for cell in ws[hdr_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1E3A8A")
        cell.alignment = Alignment(horizontal="center")
    for item in order.items.all():
        ws.append([
            item.product_name,
            "Verão" if item.line == "verao" else "Inverno" if item.line == "inverno" else "",
            item.color, item.size, item.qty, float(item.unit_price), float(item.subtotal),
        ])
    ws.append([])
    ws.append(["", "", "", "TOTAL", order.total_qty, "", float(order.total)])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    for row in ws.iter_rows(min_row=hdr_row + 1, min_col=6, max_col=7):
        for cell in row:
            cell.number_format = "#,##0.00"
    for col, width in zip("ABCDEFG", (40, 10, 18, 10, 8, 16, 16)):
        ws.column_dimensions[col].width = width
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def whatsapp_message(order: Order, sheet_url: str) -> str:
    lines = [f"Olá! Pedido *{order.code}* na Red Blue Line:", ""]
    for it in order.items.all():
        variant = " ".join(f"[{v}]" for v in (it.size, it.color) if v)
        lines.append(f"• {it.qty}x {it.product_name} {variant} — R$ {it.unit_price:.2f} = R$ {it.subtotal:.2f}")
    lines += [
        "",
        f"Peças: {order.total_qty} ({TIER_LABELS.get(order.tier, '')})",
        f"*TOTAL: R$ {order.total:.2f}*",
        "",
        f"Planilha do pedido: {sheet_url}",
    ]
    return "\n".join(lines)
