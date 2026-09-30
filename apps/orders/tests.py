import json
from decimal import Decimal
from io import BytesIO
from urllib.parse import unquote

import pytest
from openpyxl import load_workbook

from apps.orders.models import Order
from apps.orders.services import OrderError, create_order, tier_for, unit_price
from apps.products.models import Product
from tests.factories import CategoryFactory


@pytest.fixture
def products(db):
    cat = CategoryFactory(name="Camiseta UV")
    uv = Product.objects.create(name="Camiseta UV", price="59.90", category=cat, line="verao",
                                color="Branco, Azul", size_range="P ao G",
                                wholesale_price_6="49.90", is_published=True)
    jaq = Product.objects.create(name="Jaqueta", price="200.00", category=cat, line="inverno",
                                 is_published=True)
    return uv, jaq


def post(client, payload):
    return client.post("/pedido/finalizar/", data=json.dumps(payload), content_type="application/json")


def test_tiers(products):
    uv, jaq = products
    uv.refresh_from_db()
    jaq.refresh_from_db()
    assert tier_for(5) == "" and tier_for(6) == "w6" and tier_for(24) == "w24"
    assert unit_price(uv, "w6") == Decimal("49.90")
    assert unit_price(jaq, "w6") == Decimal("195.00")
    assert unit_price(jaq, "w24") == Decimal("193.00")
    uv.wholesale_price_24 = Decimal("40.00")
    assert unit_price(uv, "w24") == Decimal("40.00")
    assert unit_price(uv, "") == Decimal("59.90")


@pytest.mark.django_db
class TestCheckout:
    def test_creates_order_with_server_prices_and_whatsapp_link(self, client, products):
        uv, jaq = products
        resp = post(client, {"customer_name": "Maria", "items": [
            {"id": uv.id, "qty": 5, "size": "M", "color": "Azul", "price": 0.01},
            {"id": jaq.id, "qty": 1},
        ]})
        assert resp.status_code == 201
        data = resp.json()
        order = Order.objects.get()
        assert order.total_qty == 6 and order.tier == "w6"
        assert order.total == Decimal("5") * Decimal("49.90") + Decimal("195.00")
        assert data["total"] == f"{order.total:.2f}"
        assert data["whatsapp_url"].startswith("https://wa.me/5511954294886?text=")
        msg = unquote(data["whatsapp_url"])
        assert "RBL-" in msg and "5x Camiseta UV [M] [Azul]" in msg and "planilha.xlsx" in msg
        assert order.customer_name == "Maria"

        sheet = client.get(data["spreadsheet_url"].replace("http://testserver", ""))
        assert sheet.status_code == 200
        assert "spreadsheetml" in sheet["Content-Type"]
        ws = load_workbook(BytesIO(sheet.content)).active
        values = [c for row in ws.iter_rows(values_only=True) for c in row if c is not None]
        assert "Camiseta UV" in values and "Azul" in values and "Cliente: Maria" in values
        assert float(order.total) in values

    @pytest.mark.parametrize("payload,msg", [
        ({"items": []}, "Carrinho vazio"),
        ({"items": [{"id": "x"}]}, "Item inválido"),
        ({"items": [{"id": 99999, "qty": 1}]}, "indisponível"),
        ({"items": "abc"}, "Carrinho vazio"),
    ])
    def test_invalid_payloads(self, client, products, payload, msg):
        resp = post(client, payload)
        assert resp.status_code == 400 and msg in resp.json()["error"]

    def test_invalid_qty_size_and_json(self, client, products):
        uv, _ = products
        assert post(client, {"items": [{"id": uv.id, "qty": 0}]}).status_code == 400
        assert post(client, {"items": [{"id": uv.id, "qty": "a"}]}).status_code == 400
        assert post(client, {"items": [{"id": uv.id, "qty": 1, "size": "XXL"}]}).status_code == 400
        assert client.post("/pedido/finalizar/", data="{", content_type="application/json").status_code == 400
        assert client.get("/pedido/finalizar/").status_code == 405

    def test_unpublished_product_rejected(self, products):
        uv, _ = products
        uv.is_published = False
        uv.save()
        with pytest.raises(OrderError):
            create_order([{"id": uv.id, "qty": 1}])
        with pytest.raises(OrderError):
            create_order([{"id": 1}] * 101)

    def test_admin_lists_orders(self, admin_client, products):
        order = create_order([{"id": products[0].id, "qty": 2}])
        assert str(order) == f"Pedido #{order.pk}" and str(order.items.first()) == "2x Camiseta UV"
        resp = admin_client.get("/admin/orders/order/")
        assert resp.status_code == 200 and "Baixar XLSX" in resp.content.decode()
        assert admin_client.get(f"/admin/orders/order/{order.pk}/change/").status_code == 200
        assert admin_client.get("/admin/orders/order/add/").status_code == 200

    def test_csrf_enforced(self, products):
        from django.test import Client
        c = Client(enforce_csrf_checks=True)
        resp = c.post("/pedido/finalizar/", data=json.dumps({"items": []}), content_type="application/json")
        assert resp.status_code == 403
